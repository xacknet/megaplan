from django.shortcuts import render, redirect
from django.contrib.auth import login, authenticate, logout
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, Http404
from .forms import UserRegistrationForm, UserLoginForm
from .models import UserProfile
from functools import wraps

def role_required(role_name):
    def decorator(view_func):
        @wraps(view_func)
        def _wrapped_view(request, *args, **kwargs):
            if not request.user.is_authenticated:
                # Or redirect to login page: return redirect('login_url_name')
                raise Http404 
            try:
                if request.user.profile.role == role_name:
                    return view_func(request, *args, **kwargs)
                else:
                    # Or a custom "permission denied" page
                    return HttpResponse("You do not have the required role to view this page.", status=403)
            except UserProfile.DoesNotExist:
                 # Or a custom "profile not found" page
                return HttpResponse("User profile not found. Please contact an administrator.", status=403)
            except AttributeError: # Handles cases where user has no profile attribute
                return HttpResponse("User profile not properly configured.", status=403)
        return _wrapped_view
    return decorator

def register_view(request):
    if request.method == 'POST':
        form = UserRegistrationForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.set_password(form.cleaned_data['password'])
            user.save()
            # UserProfile signal should have created the profile. Now set the role.
            # The signal in models.py handles profile creation.
            # We just need to set the role for the newly created profile.
            try:
                # The profile is created by the signal upon User save.
                # We fetch it and assign the role.
                profile = user.profile 
                profile.role = form.cleaned_data['role']
                profile.save()
            except UserProfile.DoesNotExist:
                # This should ideally not happen if signals are working correctly
                UserProfile.objects.create(user=user, role=form.cleaned_data['role'])
            except AttributeError: # If user.profile doesn't exist for some reason
                 UserProfile.objects.create(user=user, role=form.cleaned_data['role'])


            # Log the user in directly after registration
            # Use the backend that doesn't require current password for login after registration
            # For standard username/password, the default backend works.
            new_user = authenticate(request, username=form.cleaned_data['username'], password=form.cleaned_data['password'])
            if new_user:
                login(request, new_user)
                return redirect('home') 
            else:
                # This case should be rare if user.save() and set_password() worked.
                # Add an error message or redirect to login page
                form.add_error(None, "Account created, but auto-login failed. Please try logging in manually.")


    else:
        form = UserRegistrationForm()
    return render(request, 'core/register.html', {'form': form, 'title': 'Register'})

def login_view(request):
    if request.user.is_authenticated:
        return redirect('home')
        
    if request.method == 'POST':
        form = UserLoginForm(request.POST)
        if form.is_valid():
            username = form.cleaned_data['username']
            password = form.cleaned_data['password']
            user = authenticate(request, username=username, password=password)
            if user is not None:
                login(request, user)
                next_url = request.GET.get('next')
                if next_url:
                    return redirect(next_url)
                return redirect('home')
            else:
                form.add_error(None, "Invalid username or password.")
    else:
        form = UserLoginForm()
    return render(request, 'core/login.html', {'form': form, 'title': 'Login'})

@login_required
def logout_view(request):
    logout(request)
    return redirect('login')

@login_required
@role_required('administrator')
def admin_only_view(request):
    return HttpResponse("<h1>Administrator Dashboard</h1><p>Welcome, Administrator!</p><a href=\"/\">Home</a>")

# Placeholder for 'home' view accessible to any logged-in user
@login_required
def home_view(request):
    try:
        role_display = request.user.profile.get_role_display()
    except (UserProfile.DoesNotExist, AttributeError):
        role_display = "Not assigned"
    
    context = {
        'title': 'Home',
        'role_display': role_display,
    }
    return render(request, 'core/home.html', context)

from django.contrib.gis.geos import GEOSGeometry
from django.contrib.gis.forms.fields import GeometryField
from django.core.exceptions import ValidationError as DjangoValidationError
from .forms import FieldForm
from .models import Field
import json

@login_required
def field_list_view(request):
    if request.user.profile.role == 'administrator':
        fields = Field.objects.all()
    else:
        fields = Field.objects.filter(owner=request.user)
    return render(request, 'core/field_list.html', {'fields': fields, 'title': 'My Fields'})

@login_required
def field_create_view(request): # This view now handles both create and update
    edit_id = request.GET.get('edit_id')
    instance = None
    if edit_id:
        instance = get_object_or_404(Field, pk=edit_id)
        # Permission check for editing
        if not (instance.owner == request.user or (hasattr(request.user, 'profile') and request.user.profile.role == 'administrator')):
            raise Http404("You do not have permission to edit this field.")

    if request.method == 'POST':
        form = FieldForm(request.POST, instance=instance) # Pass instance if editing
        if form.is_valid():
            try:
                geojson_str = form.cleaned_data['geometry']
                try:
                    geojson_data = json.loads(geojson_str)
                except json.JSONDecodeError:
                    form.add_error('geometry', "Invalid GeoJSON data.")
                    raise DjangoValidationError("Invalid GeoJSON data.")

                geom = GEOSGeometry(json.dumps(geojson_data.get('geometry', geojson_data))) # Handle if geojson_data is full Feature or just geometry
                
                if geom.geom_type not in ['Polygon', 'MultiPolygon']:
                    form.add_error('geometry', f"Invalid geometry type: {geom.geom_type}. Expected Polygon or MultiPolygon.")
                    raise DjangoValidationError(f"Invalid geometry type: {geom.geom_type}. Expected Polygon or MultiPolygon.")

                field = form.save(commit=False)
                if not instance: # Only set owner if creating new
                    field.owner = request.user
                field.geometry = geom
                field.save()
                return redirect('field_detail', pk=field.pk) # Redirect to detail view
            except DjangoValidationError:
                pass 
            except Exception as e:
                form.add_error(None, f"An unexpected error occurred: {str(e)}")
                
    else: # GET request
        form = FieldForm(instance=instance) # Pass instance for pre-population
    
    title = 'Edit Field' if instance else 'Create New Field'
    return render(request, 'core/field_form.html', {'form': form, 'title': title, 'field_instance': instance})


@login_required
def field_detail_view(request, pk):
    try:
        field = Field.objects.get(pk=pk)
        # Ensure the user has permission to view (owner or admin)
        if not (field.owner == request.user or request.user.profile.role == 'administrator'):
            raise Http404("You do not have permission to view this field.")
    except Field.DoesNotExist:
        raise Http404("Field not found.")
    
    # Pass the geometry as GeoJSON to the template for Leaflet
    field_geojson = field.geometry.geojson if field.geometry else None
    operations = AgroOperation.objects.filter(field=field) # Get operations for this field
    
    ndvi_tile_url = None
    mean_ndvi = None
    gee_error = None
    
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    if start_date_str and end_date_str and field_geojson:
        from . import gee_utils # Import GEE utils
        if gee_utils.initialize_gee(): # Initialize GEE
            ndvi_data = gee_utils.get_ndvi_for_field(field_geojson, start_date_str, end_date_str)
            if ndvi_data:
                if 'error' in ndvi_data:
                    gee_error = ndvi_data['error']
                else:
                    ndvi_tile_url = ndvi_data.get('tile_url')
                    mean_ndvi = ndvi_data.get('mean_ndvi')
            else:
                gee_error = "Failed to retrieve NDVI data from GEE."
        else:
            gee_error = "Failed to initialize Google Earth Engine. Check server logs and GEE configuration in settings.py."

    return render(request, 'core/field_detail.html', {
        'field': field, 
        'field_geojson': field_geojson, 
        'title': field.name,
        'ndvi_tile_url': ndvi_tile_url,
        'mean_ndvi': mean_ndvi,
        'gee_error': gee_error,
        'start_date': start_date_str,
        'end_date': end_date_str,
        'operations': operations
    })

from .forms import AgroOperationForm
from .models import AgroOperation # Crop, OperationType already implicitly available via AgroOperationForm
from django.shortcuts import get_object_or_404
from django.urls import reverse_lazy

@login_required
def agrooperation_create_view(request, field_pk):
    field = get_object_or_404(Field, pk=field_pk)
    # Permission check: only field owner or admin can add operations
    if not (field.owner == request.user or (hasattr(request.user, 'profile') and request.user.profile.role == 'administrator')):
        raise Http404("You do not have permission to add operations to this field.")

    if request.method == 'POST':
        form = AgroOperationForm(request.POST, request=request, field_instance=field)
        if form.is_valid():
            operation = form.save(commit=False)
            operation.field = field
            operation.user = request.user
            operation.save()
            return redirect('field_detail', pk=field.pk)
    else:
        form = AgroOperationForm(request=request, field_instance=field)
    
    return render(request, 'core/agrooperation_form.html', {
        'form': form, 
        'title': f'Add Operation to {field.name}',
        'field': field
    })

@login_required
def agrooperation_update_view(request, pk):
    operation = get_object_or_404(AgroOperation, pk=pk)
    field = operation.field

    # Permission check: only operation creator or admin can edit
    if not (operation.user == request.user or (hasattr(request.user, 'profile') and request.user.profile.role == 'administrator')):
        raise Http404("You do not have permission to edit this operation.")

    if request.method == 'POST':
        form = AgroOperationForm(request.POST, instance=operation, request=request, field_instance=field)
        if form.is_valid():
            form.save()
            return redirect('field_detail', pk=field.pk)
    else:
        form = AgroOperationForm(instance=operation, request=request, field_instance=field)
        
    return render(request, 'core/agrooperation_form.html', {
        'form': form, 
        'title': f'Edit Operation on {field.name}',
        'field': field,
        'operation': operation
    })

@login_required
def agrooperation_delete_view(request, pk):
    operation = get_object_or_404(AgroOperation, pk=pk)
    field = operation.field

    # Permission check: only operation creator or admin can delete
    if not (operation.user == request.user or (hasattr(request.user, 'profile') and request.user.profile.role == 'administrator')):
        raise Http404("You do not have permission to delete this operation.")

    if request.method == 'POST':
        operation.delete()
        return redirect('field_detail', pk=field.pk)
        
    return render(request, 'core/agrooperation_confirm_delete.html', {
        'operation': operation,
        'field': field,
        'title': f'Delete Operation on {field.name}'
    })

from . import analytics_utils
from .pdf_utils import generate_operations_pdf
from django.http import HttpResponse as DjangoHttpResponse # Renamed to avoid conflict
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment

@login_required
def report_dashboard_view(request):
    total_area_hectares = analytics_utils.get_total_area_for_user(request.user)
    
    if hasattr(request.user, 'profile') and request.user.profile.role == 'administrator':
        fields_for_reports = Field.objects.all().order_by('name')
    else:
        fields_for_reports = Field.objects.filter(owner=request.user).order_by('name')
        
    context = {
        'title': 'Reports Dashboard',
        'total_area_hectares': round(total_area_hectares, 2),
        'fields_for_reports': fields_for_reports,
    }
    return render(request, 'core/report_dashboard.html', context)

@login_required
def download_field_operations_report_excel(request, field_id):
    field = get_object_or_404(Field, pk=field_id)
    # Permission check
    if not (field.owner == request.user or (hasattr(request.user, 'profile') and request.user.profile.role == 'administrator')):
        raise Http404("You do not have permission to download this report.")

    operations = analytics_utils.get_operations_for_field(field)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = f"Operations_{field.name.replace(' ', '_')}"

    # Headers
    headers = ["Date", "Operation Type", "Crop", "Description", "Equipment Used", "Materials Cost (€)", "Recorded By"]
    for col_num, header_title in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num, value=header_title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        # Set column width (approximate)
        column_letter = get_column_letter(col_num)
        if header_title == "Description":
            sheet.column_dimensions[column_letter].width = 40
        elif header_title == "Operation Type":
            sheet.column_dimensions[column_letter].width = 20
        else:
            sheet.column_dimensions[column_letter].width = 15


    # Data
    for row_num, op in enumerate(operations, 2):
        sheet.cell(row=row_num, column=1, value=op.operation_date)
        sheet.cell(row=row_num, column=2, value=op.operation_type.name)
        sheet.cell(row=row_num, column=3, value=op.crop.name if op.crop else "N/A")
        sheet.cell(row=row_num, column=4, value=op.description)
        sheet.cell(row=row_num, column=5, value=op.equipment_used or "N/A")
        sheet.cell(row=row_num, column=6, value=op.materials_cost if op.materials_cost is not None else None)
        if op.materials_cost is not None: # Format as currency if value exists
             sheet.cell(row=row_num, column=6).number_format = '#,##0.00€'
        sheet.cell(row=row_num, column=7, value=op.user.username if op.user else "N/A")
        
        # Apply text wrapping for description
        sheet.cell(row=row_num, column=4).alignment = Alignment(wrap_text=True, vertical="top")


    response = DjangoHttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="operations_report_{field.name.replace(" ", "_")}.xlsx"'
    workbook.save(response)
    return response

@login_required
def download_field_operations_report_pdf(request, field_id):
    field = get_object_or_404(Field, pk=field_id)
    # Permission check
    if not (field.owner == request.user or (hasattr(request.user, 'profile') and request.user.profile.role == 'administrator')):
        raise Http404("You do not have permission to download this report.")

    operations = analytics_utils.get_operations_for_field(field)
    return generate_operations_pdf(operations, field.name)
